"""One-page ranges sheet for the pattern makers (docs/manipulations/ranges.xlsx).

Recipe, parameter, least, most, default; and each recipe's options with their default. Read from
section 3 of every draft in docs/manipulations/, so the sheet always lists what the drafts list.
The yellow cells are the pattern makers'. Nothing is filled in: a range is a drafting rule.

    uv run python scripts/make_ranges_sheet.py
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
DRAFTS = ROOT / "docs/manipulations"
OUT = DRAFTS / "ranges.xlsx"
FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F2933")
INPUT = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFC5CC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def cells(line: str) -> list[str]:
    return [c.strip().strip("`") for c in line.strip().strip("|").split("|")]


def read_draft(path: Path) -> tuple[str, list[list[str]], list[list[str]]]:
    text = path.read_text()
    name = re.search(r"Recipe name: `([a-z_]+)`", text)
    section = text[text.index("## 3.") : text.index("## 4.")]
    params: list[list[str]] = []
    options: list[list[str]] = []
    mode = ""
    for line in section.splitlines():
        if line.startswith("| Name | Meaning"):
            mode = "p"
        elif line.startswith("| Name | Choices"):
            mode = "o"
        elif line.startswith("| ---"):
            continue
        elif line.startswith("|") and mode:
            row = cells(line)
            (params if mode == "p" else options).append(row)
        elif line.strip():
            mode = ""
    return (name.group(1) if name else path.stem), params, options


def style(cell, *, bold=False, color="000000", fill=None, size=10) -> None:
    cell.font = Font(name=FONT, size=size, bold=bold, color=color)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    if fill is not None:
        cell.fill = fill


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ranges"
    ws.sheet_view.showGridLines = False
    widths = [26, 18, 46, 8, 12, 12, 12, 36]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws["A1"] = "Recipe parameters: least, most and default"
    style(ws["A1"], bold=True, size=14)
    ws.merge_cells("A2:H2")
    ws["A2"] = (
        "Fill the yellow cells. Least, most and default are numbers, or formulas over measurements when "
        "the right amount depends on size (for example hip_circ/10): they must hold at every size of the "
        "multisize table. Leave a cell empty if you do not know. The grey row is only an example of the format."
    )
    style(ws["A2"], color="52606D")
    ws.row_dimensions[2].height = 42
    head = ["Recipe", "Parameter", "What it is", "Unit", "Least", "Most", "Default", "Comment"]
    r = 4
    for c, h in enumerate(head, start=1):
        style(ws.cell(row=r, column=c, value=h), bold=True, color="FFFFFF", fill=HEAD)
    example = [
        "(example) hem_flare",
        "amount",
        "how much wider the hem gets",
        "cm",
        "0",
        "bust_circ/10",
        "2",
        "overwrite me",
    ]
    r += 1
    for c, v in enumerate(example, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        style(cell, color="9AA5B1")
        cell.border = BORDER
    drafts = [read_draft(p) for p in sorted(DRAFTS.glob("*.md")) if p.name != "TEMPLATE.md"]
    for name, params, _ in drafts:
        for p in params:
            r += 1
            vals = [name, p[0], p[1], p[2], "", "", "", ""]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=c, value=v or None)
                style(cell, fill=INPUT if c >= 5 else None)
                cell.border = BORDER
    r += 2
    ws.cell(row=r, column=1, value="Options: the choices a recipe branches on")
    style(ws.cell(row=r, column=1), bold=True, size=12)
    r += 1
    oh = ["Recipe", "Option", "Choices", "", "", "", "Default", "Comment"]
    for c, h in enumerate(oh, start=1):
        style(ws.cell(row=r, column=c, value=h or None), bold=True, color="FFFFFF", fill=HEAD)
    for name, _, options in drafts:
        for o in options:
            r += 1
            vals = [name, o[0], o[1].replace("`", ""), None, None, None, None, None]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=c, value=v)
                style(cell, fill=INPUT if c >= 7 else None)
                cell.border = BORDER
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)} ({r} rows)")


if __name__ == "__main__":
    main()
