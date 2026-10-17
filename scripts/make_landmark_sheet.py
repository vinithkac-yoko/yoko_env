"""Build the landmark sheet for the pattern makers (docs/landmarks/basic-set-landmarks.xlsx).

One tab per piece of the basic set: a picture of the piece with numbered points, and a table with one
row per numbered point or curve (where it is, how Seamly2D drew it, our proposed name and why) and
empty columns for the pattern maker's verdict. Nothing proposed here is verified.

    uv run python scripts/make_landmark_sheet.py
"""

from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from landmark_internal import internal_rows, not_drawn
from landmark_rows import piece_rows
from yoko_io.render import render_outline_png

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/landmarks/basic-set-landmarks.xlsx"
PROPOSED = ROOT / "docs/landmarks/proposed_names.yaml"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F2933")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # cells the pattern maker fills in
GUESS_FILL = PatternFill("solid", fgColor="E8F0FE")
THIN = Side(style="thin", color="BFC5CC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
COLUMNS = [
    ("No.", 6),
    ("Kind", 8),
    ("Seamly2D label", 18),
    ("How it was drawn", 38),
    ("Uses measurements", 24),
    ("x (mm)", 9),
    ("y (mm)", 9),
    ("Proposed name", 30),
    ("Also named (same object)", 34),
    ("Confidence", 11),
    ("Why we think so", 46),
    ("Verdict", 14),
    ("Correct name (if Correct)", 30),
    ("Comment", 36),
]

GLOSSARY = [
    ("CB", "centre back"),
    ("CF", "centre front"),
    ("SNP", "side neck point: where the neckline meets the shoulder seam"),
    ("BP", "bust point (the basic set does not draw one; see the bodice front's internal table)"),
    ("underarm", "the lowest point of the armhole, where it meets the side seam"),
    ("shoulder_tip", "the outer end of the shoulder seam, where the armhole begins"),
    ("side_waist / side_hip", "the point on the side seam at the waist line / hip line"),
    ("inseam", "the inside leg seam of a trouser"),
    ("cap_top", "the highest point of the sleeve cap"),
    (
        "leg_a / leg_b",
        "the two ends of a dart where its legs meet the seam or the waist line; a is the one nearer the centre line",
    ),
    (
        "apex",
        "the tip of a dart, inside the piece, not on the outline. A fish dart has two: apex_upper and apex_lower",
    ),
    (
        "cap",
        "the little bump a dart leaves in the outline of the basic piece (Seamly2D draws it as three points: leg_a, cap, leg_b)",
    ),
    (
        "waist_dart1, shoulder_dart1 ...",
        "a dart is named for where it is and numbered from the centre line outwards: waist_dart1_leg_a, waist_dart1_apex, waist_dart1_cap",
    ),
    (
        "front / back",
        "the sleeve points on the front half of the cap and the back half; the sleeve's own centre line is cap_top to wrist_center",
    ),
    (
        "bust_line / waist_line / hip_line / knee_line",
        "the level guide lines the pieces are drafted from",
    ),
    ("CB_line / CF_line", "the vertical centre-back and centre-front lines"),
]


def style(cell, *, bold=False, color="000000", fill=None, wrap=True, size=10):
    cell.font = Font(name=FONT, size=size, bold=bold, color=color)
    cell.alignment = Alignment(wrap_text=wrap, vertical="top")
    if fill is not None:
        cell.fill = fill


def readme(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 100
    rows = [
        ("Landmark sheet: basic pattern set", None, "title"),
        (
            "Why",
            "The software that drafts and checks patterns refers to points by meaningful names "
            "(bodice_front.shoulder_tip) instead of Seamly2D labels (C30). Labels are numbered in the "
            "order a point was drawn and change from pattern to pattern. We need your help to name the "
            "points that matter.",
            None,
        ),
        (
            "What is in it",
            "One tab per piece (7 pieces). Each tab has a picture of the piece and two tables. The first has "
            "one row per red numbered circle: a corner point or a curve of the piece's outline. The second, "
            "'Not on the outline', has one row per blue mark i1, i2 ...: a dart's apex, the bust point, a "
            "guide line (waist, hip, bust, knee) or a centre-front or centre-back line. Blue squares are "
            "points; blue dashed lines are lines.",
            None,
        ),
        (
            "What we ask",
            "For every row, in the yellow columns: choose a Verdict. Confirm = our proposed name is right. "
            "Correct = write the right name in 'Correct name'. Reject = this point does not need a "
            "name. Rows with no proposed name: write a name in 'Correct name' and choose Correct, or "
            "leave them empty if the point needs no name. Add anything in 'Comment'.",
            None,
        ),
        (
            "Nothing is verified",
            "All proposed names are guesses from how the point was drawn and where it sits. 'Confidence' "
            "says how sure we are: high = the formula itself names the thing (a hip point uses the hip "
            "measurement); medium = shape and position make it likely; low = a guess from shape alone. "
            "Please do not trust the names.",
            None,
        ),
        (
            "Naming",
            "<piece>.<landmark>, lower-case with underscores; common abbreviations in capitals "
            "(CB, CF, SNP). See the Glossary tab. If the usual name in your trade is different, write "
            "yours: we will use it.",
            None,
        ),
        (
            "Rows that repeat a label",
            "A curve that Seamly2D draws as one long object may appear in several rows (it is cut by the "
            "points between them). Name each row for the part of the outline it shows.",
            None,
        ),
        (
            "Also named",
            "When one Seamly2D object carries two names (a point shared by two pieces, such as the "
            "underarm, or one long curve cut into several rows, such as the sleeve cap, or one guide "
            "line used by two pieces), every row that shares it says so in 'Also named'. Both names are "
            "kept: please check that each is right for its piece.",
            None,
        ),
        (
            "Darts",
            "Every dart is named the same way: leg_a, leg_b, apex and cap. leg_a and leg_b are its two "
            "ends where the legs meet the seam or waist line; apex is its tip inside the piece; cap is the "
            "bump it leaves in the outline of the basic piece. Darts are called waist_dart1, waist_dart2 ... "
            "or shoulder_dart1, counted from the centre line outwards. A fish dart has two tips: "
            "apex_upper and apex_lower.",
            None,
        ),
        (
            "Front and back",
            "The sleeve's points are named front or back. We worked out which half is which from the "
            "armhole lengths the formulas use; please confirm.",
            None,
        ),
        ("Asked for, but not drawn", None, "notdrawn"),
        (
            "Positions",
            "x and y are millimetres on the drawing at the base size (UK 6, height 164 cm). y grows "
            "downward, as on screen. They are for orientation only.",
            None,
        ),
        ("Example row", None, "example"),
    ]
    r = 2
    for head, text, kind in rows:
        if kind == "title":
            ws.cell(row=r, column=2, value=head)
            style(ws.cell(row=r, column=2), bold=True, size=16, wrap=False)
            ws.row_dimensions[r].height = 28
        elif kind == "notdrawn":
            ws.cell(row=r, column=2, value=head)
            style(ws.cell(row=r, column=2), bold=True)
            lines = not_drawn()
            ws.cell(row=r, column=3, value="\n".join(f"- {x}" for x in lines))
            style(ws.cell(row=r, column=3))
            ws.row_dimensions[r].height = max(30, 15 * (sum(len(x) // 90 + 1 for x in lines) + 1))
        elif kind == "example":
            ws.cell(row=r, column=2, value=head)
            style(ws.cell(row=r, column=2), bold=True)
            r += 1
            for c, (h, _w) in enumerate(COLUMNS, start=1):
                cell = ws.cell(row=r, column=c + 1, value=h)
                style(cell, bold=True, color="FFFFFF", fill=HEAD_FILL)
            r += 1
            example = [
                "7",
                "point",
                "C30",
                "pointOfContact; radius=shoulder_length",
                "shoulder_length",
                "329.1",
                "74.2",
                "bodice_front.shoulder_tip",
                "",
                "medium",
                "end of the shoulder; radius is shoulder_length",
                "Confirm",
                "",
                "",
            ]
            for c, v in enumerate(example, start=1):
                cell = ws.cell(row=r, column=c + 1, value=v)
                style(cell, fill=INPUT_FILL if c >= 12 else None)
            ws.cell(row=r + 1, column=2, value="(This row is only an illustration of the format.)")
            style(ws.cell(row=r + 1, column=2), color="52606D", wrap=False)
            r += 1
        else:
            ws.cell(row=r, column=2, value=head)
            style(ws.cell(row=r, column=2), bold=True)
            ws.cell(row=r, column=3, value=text)
            style(ws.cell(row=r, column=3))
            ws.row_dimensions[r].height = max(30, 15 * (len(text or "") // 95 + 1))
        r += 1
    # the example table is wider than the columns used above: give them widths
    for c in range(4, len(COLUMNS) + 3):
        ws.column_dimensions[get_column_letter(c)].width = 20


def glossary(ws, names: dict[str, list[tuple[str, str]]]) -> None:
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90
    ws["A1"] = "Abbreviations"
    style(ws["A1"], bold=True, size=14, wrap=False)
    r = 3
    for term, meaning in GLOSSARY:
        ws.cell(row=r, column=1, value=term)
        ws.cell(row=r, column=2, value=meaning)
        style(ws.cell(row=r, column=1), bold=True)
        style(ws.cell(row=r, column=2))
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="All proposed names")
    style(ws.cell(row=r, column=1), bold=True, size=14, wrap=False)
    r += 2
    for piece, entries in names.items():
        ws.cell(row=r, column=1, value=piece)
        style(ws.cell(row=r, column=1), bold=True, fill=GUESS_FILL)
        ws.cell(row=r, column=2, value=", ".join(n for n, _ in entries))
        style(ws.cell(row=r, column=2), fill=GUESS_FILL)
        ws.row_dimensions[r].height = max(30, 14 * (len(ws.cell(row=r, column=2).value) // 85 + 1))
        r += 1


def also_text(entries: list[tuple[str, str, str]], piece: str, number: str) -> str:
    """The other names of the same Seamly2D object, one per line, for the 'Also named' column."""
    lines = [
        f"{'' if p == piece else p + ' '}No. {n}: {name or '(no proposal yet)'}"
        for p, n, name in entries
        if (p, n) != (piece, number)
    ]
    return "\n".join(lines)


def table_row(ws, r: int, values: list, dv: DataValidation) -> None:
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        style(cell, fill=INPUT_FILL if c >= 12 else (GUESS_FILL if c in (8, 9, 10, 11) else None))
        cell.border = BORDER
    dv.add(ws.cell(row=r, column=12))


def table_head(ws, r: int) -> None:
    for c, (h, _w) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=r, column=c, value=h)
        style(cell, bold=True, color="FFFFFF", fill=HEAD_FILL)
        cell.border = BORDER


def piece_tab(
    ws,
    pr,
    proposals: dict[int, dict],
    internal: list,
    shared: dict[int, list[tuple[str, str, str]]],
    groups: dict[str, list[tuple[str, str, str]]],
) -> int:
    ws.sheet_view.showGridLines = False
    for c, (_h, w) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws["A1"] = pr.name
    style(ws["A1"], bold=True, size=14, wrap=False)
    ws["A2"] = (
        "Red numbered circles on the picture match the 'No.' column of the first table. Blue marks "
        "(i1, i2 ...) are the second table: things that are not on the outline. Fill the yellow "
        "columns (Verdict, Correct name, Comment)."
    )
    style(ws["A2"], color="52606D", wrap=False)
    inner_marks = [(r.number, r.points_px[0]) for r in internal if r.kind == "point"]
    inner_lines = [(r.number, r.points_px[0], r.points_px[1]) for r in internal if r.kind == "line"]
    png = render_outline_png(
        list(pr.outline_px),
        [(str(n), pt) for n, pt in pr.marks_px],
        width_px=900,
        inner_marks=inner_marks,
        inner_lines=inner_lines,
    )
    img = XlImage(BytesIO(png))
    scale = min(1.0, 620 / img.height, 900 / img.width)
    img.width, img.height = int(img.width * scale), int(img.height * scale)
    ws.add_image(img, "A4")
    header_row = 4 + int(img.height / 20) + 3
    ws.cell(row=header_row - 1, column=1, value="On the outline")
    style(ws.cell(row=header_row - 1, column=1), bold=True, size=12, wrap=False)
    table_head(ws, header_row)
    dv = DataValidation(type="list", formula1='"Confirm,Correct,Reject"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, row in enumerate(pr.rows, start=1):
        r = header_row + i
        prop = proposals.get(row.number)
        also = also_text(shared.get(row.object_id, []), pr.name, str(row.number))
        values = [
            row.number,
            row.kind,
            row.label,
            row.tool + ("; " + row.how if row.how else ""),
            row.measurements,
            row.x_mm,
            row.y_mm,
            prop["name"] if prop else "",
            also,
            prop["confidence"] if prop else "",
            prop["reason"] if prop else "no proposal: please name it if it matters",
            None,
            None,
            None,
        ]
        table_row(ws, r, values, dv)
        if also:
            ws.row_dimensions[r].height = 15 * (also.count("\n") + 1) + 10
    r = header_row + len(pr.rows) + 3
    ws.cell(row=r - 1, column=1, value="Not on the outline (blue marks on the picture)")
    style(ws.cell(row=r - 1, column=1), bold=True, size=12, wrap=False)
    if not internal:
        ws.cell(row=r, column=1, value="Nothing proposed for this piece.")
        style(ws.cell(row=r, column=1), color="52606D", wrap=False)
        return header_row
    table_head(ws, r)
    for i, row in enumerate(internal, start=1):
        also = also_text(groups.get(row.group, []), pr.name, row.number) if row.group else ""
        values = [
            row.number,
            row.kind,
            row.label,
            row.how,
            row.measurements,
            row.x_mm,
            row.y_mm,
            row.name,
            also,
            row.confidence,
            row.reason,
            None,
            None,
            None,
        ]
        table_row(ws, r + i, values, dv)
        if also:
            ws.row_dimensions[r + i].height = 15 * (also.count("\n") + 1) + 10
    return header_row


def seam_pairs(ws, piece_names: list[str]) -> None:
    """A tab for the pattern makers to say which edges are sewn to which."""
    ws.sheet_view.showGridLines = False
    widths = [22, 12, 12, 22, 12, 12, 16, 40]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws["A1"] = "Seam pairs: which edges are sewn together"
    style(ws["A1"], bold=True, size=14, wrap=False)
    ws["A2"] = (
        "One row per seam. Use the numbers from the piece tabs ('No.' column): an edge runs from one "
        "numbered corner to the next. Say if the two edges should be the same length, or by how much "
        "one is longer (ease). Leave the yellow cells empty for seams you are not sure about."
    )
    style(ws["A2"], color="52606D")
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 48
    heads = [
        "First piece", "Edge from No.", "Edge to No.", "Second piece", "Edge from No.", "Edge to No.",
        "Same length?", "Comment (ease, notches, anything)",
    ]  # fmt: skip
    for c, h in enumerate(heads, start=1):
        style(ws.cell(row=4, column=c, value=h), bold=True, color="FFFFFF", fill=HEAD_FILL)
    example = [
        "C - Bodice Front",
        12,
        15,
        "C - Bodice Back",
        9,
        6,
        "Yes",
        "Example row: please overwrite",
    ]
    for c, v in enumerate(example, start=1):
        cell = ws.cell(row=5, column=c, value=v)
        style(cell, color="9AA5B1")
        cell.border = BORDER
    pieces = DataValidation(
        type="list", formula1='"' + ",".join(piece_names) + '"', allow_blank=True
    )
    same = DataValidation(type="list", formula1='"Yes,No,Ease"', allow_blank=True)
    ws.add_data_validation(pieces)
    ws.add_data_validation(same)
    for r in range(6, 36):
        for c in range(1, 9):
            cell = ws.cell(row=r, column=c)
            style(cell, fill=INPUT_FILL)
            cell.border = BORDER
        pieces.add(ws.cell(row=r, column=1))
        pieces.add(ws.cell(row=r, column=4))
        same.add(ws.cell(row=r, column=7))


def main() -> None:
    proposed = yaml.safe_load(PROPOSED.read_text())["pieces"]
    internal = internal_rows()
    pieces = piece_rows()
    # every name a Seamly2D object has, over all pieces: object id -> (piece, No., name)
    shared: dict[int, list[tuple[str, str, str]]] = {}
    for pr in pieces:
        entries = proposed.get(pr.name, {})
        for row in pr.rows:
            name = entries.get(row.number, {}).get("name", "")
            shared.setdefault(row.object_id, []).append((pr.name, str(row.number), name))
    groups: dict[str, list[tuple[str, str, str]]] = {}
    for piece, rows in internal.items():
        for ir in rows:
            if ir.group:
                groups.setdefault(ir.group, []).append((piece, ir.number, ir.name))
    wb = Workbook()
    readme(wb.active)
    wb.active.title = "Read me"
    names: dict[str, list[tuple[str, str]]] = {}
    for pr in pieces:
        entries = proposed.get(pr.name, {})
        for number, e in entries.items():
            row = next((r for r in pr.rows if r.number == number), None)
            if row is None or row.label != e["label"]:
                raise SystemExit(
                    f"{pr.name} row {number}: proposal is for {e['label']}, "
                    f"the piece has {row.label if row else 'nothing'} there"
                )
        extra = [(r.name, r.confidence) for r in internal.get(pr.name, [])]
        names[pr.name] = [(e["name"], e["confidence"]) for _, e in sorted(entries.items())] + extra
        ws = wb.create_sheet(pr.name[:31])
        piece_tab(ws, pr, entries, internal.get(pr.name, []), shared, groups)
    gl = wb.create_sheet("Glossary", 1)
    glossary(gl, names)
    seam_pairs(wb.create_sheet("Seam pairs", 2), list(names))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
